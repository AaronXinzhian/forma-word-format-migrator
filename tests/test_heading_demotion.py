#!/usr/bin/env python3
"""Regression coverage for the optional one-level heading demotion mode.

[INPUT]: 依赖 __future__, contextlib, io, inspect, sys, tempfile, unittest, zipfile, pathlib, typing, docx, docx.oxml, docx.oxml.ns, docx.shared, lxml, style_pack_manager, word_style_transfer
[OUTPUT]: 验证样式与直接大纲标题映射、编号、单次降级及正文和表格保护
[POS]: 标题格式回归测试，通过真实格式包应用与直接迁移验证内容语义
[PROTOCOL]: 变更时更新此头部，然后检查上级 FOLDER_INDEX.md
"""

from __future__ import annotations

import contextlib
import io
import inspect
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path
from typing import Dict, Iterable, Tuple
from unittest import mock

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn as docx_qn
from docx.shared import Pt, RGBColor, Twips
from lxml import etree


TEST_DIR = Path(__file__).resolve().parent
PROJECT_DIR = TEST_DIR.parent
sys.path.insert(0, str(PROJECT_DIR))

import style_pack_manager as manager  # noqa: E402
import word_style_transfer as core  # noqa: E402


NS = {"w": core.W_NS}
BODY_HEADINGS: Tuple[Tuple[int, str], ...] = tuple(
    (level, "目标%d级标题" % level) for level in range(1, 6)
)
DEEPEST_HEADING = "目标9级标题"
HEADER_HEADING = "页眉标题不降级"
SOURCE_LEFT_BY_LEVEL = {
    level: 240 + level * 137 for level in range(1, 10)
}


def qn(local: str) -> str:
    return "{%s}%s" % (core.W_NS, local)


def _value_child(tag: str, value: object) -> etree._Element:
    child = OxmlElement(tag)
    child.set(docx_qn("w:val"), str(value))
    return child


def _next_numbering_id(root: etree._Element, attribute: str) -> int:
    values = []
    for child in root:
        raw = child.get(docx_qn(attribute))
        if raw is None:
            continue
        try:
            values.append(int(raw))
        except ValueError:
            continue
    return max(values, default=-1) + 1


def _set_style_numbering(style: object, num_id: int, level: int) -> None:
    num_pr = style._element.get_or_add_pPr().get_or_add_numPr()
    num_pr.get_or_add_ilvl().val = level
    num_pr.get_or_add_numId().val = num_id


def _add_nine_level_heading_numbering(doc: Document) -> int:
    """Create one genuine H1-H9 Word outline numbering definition."""
    numbering = doc.part.numbering_part.element
    abstract_id = _next_numbering_id(numbering, "w:abstractNumId")
    num_id = _next_numbering_id(numbering, "w:numId")

    abstract = OxmlElement("w:abstractNum")
    abstract.set(docx_qn("w:abstractNumId"), str(abstract_id))
    abstract.append(_value_child("w:nsid", "D3A07E09"))
    abstract.append(_value_child("w:multiLevelType", "multilevel"))
    abstract.append(_value_child("w:tmpl", "D3A07E99"))

    for level in range(9):
        lvl = OxmlElement("w:lvl")
        lvl.set(docx_qn("w:ilvl"), str(level))
        lvl.append(_value_child("w:start", 1))
        lvl.append(_value_child("w:numFmt", "decimal"))
        lvl.append(_value_child("w:pStyle", "Heading%d" % (level + 1)))
        lvl.append(_value_child("w:suff", "tab"))
        lvl.append(
            _value_child(
                "w:lvlText",
                ".".join("%%%d" % component for component in range(1, level + 2)),
            )
        )
        lvl.append(_value_child("w:lvlJc", "left"))

        position = 480 + level * 240
        paragraph_properties = OxmlElement("w:pPr")
        tabs = OxmlElement("w:tabs")
        tab = OxmlElement("w:tab")
        tab.set(docx_qn("w:val"), "num")
        tab.set(docx_qn("w:pos"), str(position))
        tabs.append(tab)
        paragraph_properties.append(tabs)
        indentation = OxmlElement("w:ind")
        indentation.set(docx_qn("w:left"), str(position))
        indentation.set(docx_qn("w:hanging"), "180")
        paragraph_properties.append(indentation)
        lvl.append(paragraph_properties)
        abstract.append(lvl)

    insertion = len(numbering)
    for index, child in enumerate(numbering):
        if child.tag == docx_qn("w:num"):
            insertion = index
            break
    numbering.insert(insertion, abstract)

    concrete = OxmlElement("w:num")
    concrete.set(docx_qn("w:numId"), str(num_id))
    concrete.append(_value_child("w:abstractNumId", abstract_id))
    numbering.append(concrete)

    for level in range(9):
        _set_style_numbering(doc.styles["Heading %d" % (level + 1)], num_id, level)
    return num_id


def _make_nine_level_numbered_source(path: Path) -> None:
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10)

    for level in range(1, 10):
        style = doc.styles["Heading %d" % level]
        style.font.name = "Arial"
        style.font.size = Pt(max(9, 22 - level))
        style.font.bold = True
        style.font.color.rgb = RGBColor(0x20, 0x40 + level, 0x60)
        paragraph_properties = style._element.get_or_add_pPr()
        outline = paragraph_properties.find(docx_qn("w:outlineLvl"))
        if outline is None:
            outline = OxmlElement("w:outlineLvl")
            paragraph_properties.append(outline)
        outline.set(docx_qn("w:val"), str(level - 1))

    _add_nine_level_heading_numbering(doc)
    for level in range(1, 10):
        paragraph = doc.add_heading("格式源%d级标题" % level, level=level)
        paragraph.paragraph_format.left_indent = Twips(
            SOURCE_LEFT_BY_LEVEL[level]
        )
    doc.add_paragraph("格式源正文。")
    doc.save(path)


def _make_three_level_numbered_source(path: Path) -> None:
    """Create a source that can infer deeper logic but only uses H1-H3."""
    doc = Document()
    doc.styles["Normal"].font.size = Pt(10)
    for level in range(1, 4):
        style = doc.styles["Heading %d" % level]
        style.font.name = "Arial"
        style.font.size = Pt(22 - level)
        style.font.bold = True
    _add_nine_level_heading_numbering(doc)
    for level in range(1, 4):
        paragraph = doc.add_heading("三级源%d级标题" % level, level=level)
        paragraph.paragraph_format.left_indent = Twips(
            SOURCE_LEFT_BY_LEVEL[level]
        )
    doc.save(path)


def _make_plain_demotion_target(path: Path) -> None:
    doc = Document()
    for level, text in BODY_HEADINGS:
        paragraph = doc.add_heading(text, level=level)
        paragraph.paragraph_format.left_indent = Twips(7000 + level)
    deepest = doc.add_heading(DEEPEST_HEADING, level=9)
    deepest.paragraph_format.left_indent = Twips(7999)
    doc.add_paragraph("正文中的 1.2、2026 和标题文字都必须保留。")

    header = doc.sections[0].header
    header_paragraph = header.paragraphs[0]
    header_paragraph.style = doc.styles["Heading 1"]
    header_paragraph.text = HEADER_HEADING
    doc.save(path)


def _make_prefixed_demotion_target(path: Path) -> None:
    doc = Document()
    prefixes: Iterable[Tuple[int, str, str]] = (
        (1, "1  ", "手工一级"),
        (2, "1.1  ", "手工二级"),
        (3, "1.1.1  ", "手工三级"),
        (4, "1.1.1.1  ", "手工四级"),
        (5, "1.1.1.1.1  ", "手工五级"),
    )
    for level, prefix, title in prefixes:
        paragraph = doc.add_paragraph(style="Heading %d" % level)
        # Split several prefixes across runs to keep the cleanup path honest.
        midpoint = max(1, len(prefix) // 2)
        paragraph.add_run(prefix[:midpoint])
        paragraph.add_run(prefix[midpoint:])
        paragraph.add_run(title)
    doc.add_paragraph("正文保留 1.1.1 和 2026 数字。")
    doc.save(path)


def _set_direct_outline(paragraph: object, value: object) -> None:
    properties = paragraph._element.get_or_add_pPr()
    outline = properties.find(docx_qn("w:outlineLvl"))
    if outline is None:
        outline = OxmlElement("w:outlineLvl")
        properties.append(outline)
    outline.set(docx_qn("w:val"), str(value))


def _set_explicit_normal_style(paragraph: object) -> None:
    # python-docx removes pStyle when assigning its default Normal style.
    # Insert it explicitly so this fixture exercises Normal + outlineLvl.
    properties = paragraph._element.get_or_add_pPr()
    style = properties.find(docx_qn("w:pStyle"))
    if style is None:
        style = OxmlElement("w:pStyle")
        properties.insert(0, style)
    style.set(docx_qn("w:val"), "Normal")


def read_zip(path: Path) -> Dict[str, bytes]:
    with zipfile.ZipFile(path) as archive:
        bad = archive.testzip()
        if bad:
            raise AssertionError("damaged package member: %s" % bad)
        return {name: archive.read(name) for name in archive.namelist()}


def xml(entries: Dict[str, bytes], name: str) -> etree._Element:
    return etree.fromstring(entries[name])


def paragraph_for_text(root: etree._Element, text: str) -> etree._Element:
    for paragraph in root.xpath("//w:p", namespaces=NS):
        value = "".join(paragraph.xpath(".//w:t/text()", namespaces=NS))
        if value == text:
            return paragraph
    raise AssertionError("paragraph not found: %s" % text)


def paragraph_style(paragraph: etree._Element) -> str:
    style = paragraph.find("w:pPr/w:pStyle", namespaces=NS)
    if style is None:
        raise AssertionError("paragraph has no style")
    return str(style.get(qn("val")))


def direct_numbering(paragraph: etree._Element) -> Tuple[str, int]:
    num_id = paragraph.find("w:pPr/w:numPr/w:numId", namespaces=NS)
    level = paragraph.find("w:pPr/w:numPr/w:ilvl", namespaces=NS)
    if num_id is None or level is None:
        raise AssertionError("paragraph has no complete direct numbering")
    return str(num_id.get(qn("val"))), int(level.get(qn("val")))


def direct_left_indent(paragraph: etree._Element) -> int:
    indentation = paragraph.find("w:pPr/w:ind", namespaces=NS)
    if indentation is None or indentation.get(qn("left")) is None:
        raise AssertionError("paragraph has no direct left indent")
    return int(indentation.get(qn("left")))


class HeadingDemotionTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        cls._temporary = tempfile.TemporaryDirectory(
            prefix="word-heading-demotion-tests-"
        )
        cls.working_dir = Path(cls._temporary.name)
        cls.source = cls.working_dir / "nine-level-source.docx"
        cls.target = cls.working_dir / "plain-target.docx"
        cls.prefixed_target = cls.working_dir / "prefixed-target.docx"
        cls.pack = cls.working_dir / "nine-level.wfstyle"
        _make_nine_level_numbered_source(cls.source)
        _make_plain_demotion_target(cls.target)
        _make_prefixed_demotion_target(cls.prefixed_target)
        cls.manifest = manager.create_style_pack(
            cls.source, cls.pack, display_name="九级标题降级测试"
        )

    @classmethod
    def tearDownClass(cls) -> None:
        cls._temporary.cleanup()

    def _apply(
        self, output_name: str, *, demote_headings: bool = False, prefixed: bool = False
    ) -> Tuple[Dict[str, bytes], core.TransferStats]:
        output = self.working_dir / output_name
        arguments = {}
        if demote_headings:
            arguments["demote_headings"] = True
        _manifest, stats = manager.apply_style_pack(
            self.pack,
            self.prefixed_target if prefixed else self.target,
            output,
            **arguments,
        )
        return read_zip(output), stats

    def test_public_interfaces_expose_opt_in_demotion_with_safe_defaults(self) -> None:
        manager_parameter = inspect.signature(manager.apply_style_pack).parameters[
            "demote_headings"
        ]
        self.assertIs(manager_parameter.default, False)
        core_parameter = inspect.signature(core.clean_content_xml).parameters[
            "heading_level_shift"
        ]
        self.assertEqual(core_parameter.default, 0)

        parser = manager.build_parser()
        disabled = parser.parse_args(
            ["apply-pack", "--pack", "a", "--target", "b", "--out", "c"]
        )
        enabled = parser.parse_args(
            [
                "apply-pack",
                "--pack",
                "a",
                "--target",
                "b",
                "--out",
                "c",
                "--demote-headings",
            ]
        )
        self.assertIs(disabled.demote_headings, False)
        self.assertIs(enabled.demote_headings, True)

        with mock.patch.object(
            manager,
            "apply_style_pack",
            return_value=({}, core.TransferStats()),
        ) as apply_mock, contextlib.redirect_stdout(io.StringIO()):
            exit_code = manager.main(
                [
                    "apply-pack",
                    "--pack",
                    "a",
                    "--target",
                    "b.docx",
                    "--out",
                    "c.docx",
                    "--demote-headings",
                ]
            )
        self.assertEqual(exit_code, 0)
        self.assertIs(apply_mock.call_args.kwargs["demote_headings"], True)

    def test_default_apply_keeps_heading_levels_and_all_body_text(self) -> None:
        entries, _stats = self._apply("default.docx")
        target_entries = read_zip(self.target)
        document = xml(entries, "word/document.xml")
        self.assertEqual(
            document.xpath("//w:t/text()", namespaces=NS),
            xml(target_entries, "word/document.xml").xpath(
                "//w:t/text()", namespaces=NS
            ),
        )

        num_ids = set()
        for source_level, text in BODY_HEADINGS:
            paragraph = paragraph_for_text(document, text)
            self.assertEqual(paragraph_style(paragraph), "Heading%d" % source_level)
            num_id, number_level = direct_numbering(paragraph)
            num_ids.add(num_id)
            self.assertEqual(number_level, source_level - 1)
            self.assertEqual(
                direct_left_indent(paragraph), SOURCE_LEFT_BY_LEVEL[source_level]
            )
        self.assertEqual(len(num_ids), 1)

    def test_enabled_apply_demotes_h1_through_h5_to_h2_through_h6(self) -> None:
        entries, _stats = self._apply(
            "demoted.docx", demote_headings=True
        )
        target_entries = read_zip(self.target)
        document = xml(entries, "word/document.xml")
        self.assertEqual(
            document.xpath("//w:t/text()", namespaces=NS),
            xml(target_entries, "word/document.xml").xpath(
                "//w:t/text()", namespaces=NS
            ),
        )

        num_ids = set()
        for source_level, text in BODY_HEADINGS:
            destination_level = source_level + 1
            paragraph = paragraph_for_text(document, text)
            self.assertEqual(
                paragraph_style(paragraph), "Heading%d" % destination_level
            )
            num_id, number_level = direct_numbering(paragraph)
            num_ids.add(num_id)
            self.assertEqual(number_level, destination_level - 1)
            self.assertEqual(
                direct_left_indent(paragraph),
                SOURCE_LEFT_BY_LEVEL[destination_level],
            )
        self.assertEqual(len(num_ids), 1)

    def test_h9_stays_h9_warns_and_header_heading_is_not_demoted(self) -> None:
        entries, stats = self._apply(
            "demoted-boundary.docx", demote_headings=True
        )
        document = xml(entries, "word/document.xml")
        deepest = paragraph_for_text(document, DEEPEST_HEADING)
        self.assertEqual(paragraph_style(deepest), "Heading9")
        self.assertEqual(direct_numbering(deepest)[1], 8)
        self.assertTrue(
            any("标题" in warning and "9" in warning for warning in stats.warnings),
            stats.warnings,
        )

        header_names = sorted(
            name
            for name in entries
            if name.startswith("word/header") and name.endswith(".xml")
        )
        self.assertTrue(header_names)
        header = xml(entries, header_names[0])
        header_heading = paragraph_for_text(header, HEADER_HEADING)
        self.assertEqual(paragraph_style(header_heading), "Heading1")
        self.assertEqual(direct_numbering(header_heading)[1], 0)

    def test_manager_scopes_heading_shift_to_the_main_document_story(self) -> None:
        original = core.clean_content_xml
        signature = inspect.signature(original)
        with mock.patch.object(core, "clean_content_xml", wraps=original) as cleaner:
            self._apply("demoted-scope.docx", demote_headings=True)

        shifts = {}
        for call in cleaner.call_args_list:
            bound = signature.bind_partial(*call.args, **call.kwargs)
            label = str(bound.arguments["label"])
            shifts[label] = int(bound.arguments.get("heading_level_shift", 0))
        self.assertEqual(shifts.get("word/document.xml"), 1)
        self.assertTrue(
            any(name.startswith("word/header") for name in shifts), shifts
        )
        self.assertTrue(
            all(
                shift == 0
                for name, shift in shifts.items()
                if name != "word/document.xml"
            ),
            shifts,
        )

    def test_manual_prefixes_are_cleaned_using_original_heading_levels(self) -> None:
        entries, stats = self._apply(
            "demoted-prefixed.docx", demote_headings=True, prefixed=True
        )
        document = xml(entries, "word/document.xml")
        expected_titles = tuple("手工%s级" % label for label in "一二三四五")
        actual_text = document.xpath("//w:t/text()", namespaces=NS)
        for title in expected_titles:
            self.assertIn(title, actual_text)
        self.assertIn("正文保留 1.1.1 和 2026 数字。", actual_text)
        self.assertFalse(any(text.startswith("1") for text in actual_text[:-1]))
        self.assertEqual(stats.heading_prefixes_removed, 5)

        for destination_level, title in enumerate(expected_titles, start=2):
            paragraph = paragraph_for_text(document, title)
            self.assertEqual(
                paragraph_style(paragraph), "Heading%d" % destination_level
            )
            self.assertEqual(direct_numbering(paragraph)[1], destination_level - 1)

    def test_three_level_pack_is_extended_in_memory_through_heading_nine(self) -> None:
        source = self.working_dir / "three-level-source.docx"
        pack = self.working_dir / "three-level.wfstyle"
        output = self.working_dir / "three-level-demoted.docx"
        _make_three_level_numbered_source(source)
        manifest = manager.create_style_pack(source, pack)
        self.assertEqual(
            set(map(int, manifest["heading_authorities"])),
            {0, 1, 2, 3, 4},
        )

        manager.apply_style_pack(
            pack,
            self.target,
            output,
            demote_headings=True,
        )
        entries = read_zip(output)
        document = xml(entries, "word/document.xml")
        former_h5 = paragraph_for_text(document, "目标5级标题")
        deepest = paragraph_for_text(document, DEEPEST_HEADING)
        self.assertEqual(paragraph_style(former_h5), "Heading6")
        self.assertEqual(direct_numbering(former_h5)[1], 5)
        self.assertEqual(paragraph_style(deepest), "Heading9")
        self.assertEqual(direct_numbering(deepest)[1], 8)

        styles = xml(entries, "word/styles.xml")
        for level in range(6, 10):
            self.assertEqual(
                len(
                    styles.xpath(
                        "./w:style[@w:styleId=$style_id]",
                        namespaces=NS,
                        style_id="Heading%d" % level,
                    )
                ),
                1,
            )

    def test_direct_outline_headings_receive_source_styles_and_numbering(self) -> None:
        target = self.working_dir / "outline-only-target.docx"
        output = self.working_dir / "outline-only-output.docx"
        document = Document()
        expected = []
        for level in range(9):
            for explicit_normal in (False, True):
                text = "大纲%d级-%s" % (level + 1, explicit_normal)
                paragraph = document.add_paragraph(text)
                if explicit_normal:
                    _set_explicit_normal_style(paragraph)
                _set_direct_outline(paragraph, level)
                expected.append((text, level))
        overridden = document.add_paragraph("直接大纲优先", style="Heading 1")
        _set_direct_outline(overridden, 2)
        expected.append(("直接大纲优先", 2))
        document.add_paragraph("普通正文和数字 1.2 保留。")
        document.save(target)

        _manifest, stats = manager.apply_style_pack(self.pack, target, output)
        entries = read_zip(output)
        root = xml(entries, "word/document.xml")
        catalog = core.build_style_catalog(entries["word/styles.xml"], entries["word/document.xml"])
        self.assertEqual(
            root.xpath("//w:t/text()", namespaces=NS),
            xml(read_zip(target), "word/document.xml").xpath("//w:t/text()", namespaces=NS),
        )
        for text, level in expected:
            paragraph = paragraph_for_text(root, text)
            style_id = paragraph_style(paragraph)
            self.assertEqual(style_id, "Heading%d" % (level + 1))
            self.assertEqual(catalog.resolved_outline[style_id], level)
            self.assertEqual(direct_numbering(paragraph)[1], level)
            self.assertIsNone(paragraph.find("w:pPr/w:outlineLvl", namespaces=NS))
        self.assertEqual(stats.heading_numbers_applied, len(expected))

    def test_body_outline_nine_and_invalid_values_do_not_become_headings(self) -> None:
        target = self.working_dir / "body-outline-target.docx"
        output = self.working_dir / "body-outline-output.docx"
        document = Document()
        texts = []
        for style_name in (None, "Normal", "Heading 1"):
            text = "正文大纲9-%s" % style_name
            paragraph = document.add_paragraph(text, style=style_name)
            if style_name == "Normal":
                _set_explicit_normal_style(paragraph)
            _set_direct_outline(paragraph, 9)
            texts.append(text)
        for value in (-1, 10, "invalid"):
            text = "非法大纲-%s" % value
            paragraph = document.add_paragraph(text)
            _set_direct_outline(paragraph, value)
            texts.append(text)
        visual = document.add_paragraph("只有加粗外观的正文")
        visual.runs[0].bold = True
        texts.append(visual.text)
        document.save(target)

        _manifest, stats = manager.apply_style_pack(
            self.pack, target, output, demote_headings=True
        )
        entries = read_zip(output)
        root = xml(entries, "word/document.xml")
        catalog = core.build_style_catalog(entries["word/styles.xml"], entries["word/document.xml"])
        self.assertEqual(root.xpath("//w:t/text()", namespaces=NS), texts)
        for text in texts:
            paragraph = paragraph_for_text(root, text)
            style = paragraph.find("w:pPr/w:pStyle", namespaces=NS)
            style_id = str(style.get(qn("val"))) if style is not None else catalog.fallback("paragraph")
            self.assertIsNone(catalog.resolved_outline.get(style_id))
            self.assertIsNone(paragraph.find("w:pPr/w:numPr", namespaces=NS))
        self.assertEqual(stats.heading_numbers_applied, 0)
        self.assertEqual(stats.heading_levels_demoted, 0)

    def test_explicit_body_outline_wins_over_same_id_source_heading(self) -> None:
        source = self.working_dir / "same-id-heading-source.docx"
        target = self.working_dir / "same-id-body-target.docx"
        pack = self.working_dir / "same-id-body.wfstyle"
        output = self.working_dir / "same-id-body-output.docx"
        document = Document()
        style = document.styles.add_style("ClauseBody", 1)
        outline = OxmlElement("w:outlineLvl")
        outline.set(docx_qn("w:val"), "0")
        style._element.get_or_add_pPr().append(outline)
        document.add_paragraph("模板正式一级标题", style="ClauseBody")
        document.save(source)
        document = Document()
        document.styles.add_style("ClauseBody", 1)
        paragraph = document.add_paragraph("明确正文，不得提升成标题", style="ClauseBody")
        _set_direct_outline(paragraph, 9)
        document.save(target)
        manager.create_style_pack(source, pack)
        _manifest, report = manager.preflight_style_pack(pack, target)
        self.assertEqual(report["heading_level_counts"], {})
        _manifest, stats = manager.apply_style_pack(pack, target, output)
        entries = read_zip(output)
        root = xml(entries, "word/document.xml")
        paragraph = paragraph_for_text(root, "明确正文，不得提升成标题")
        catalog = core.build_style_catalog(entries["word/styles.xml"], entries["word/document.xml"])
        self.assertIsNone(core.paragraph_outline_level(paragraph.find("w:pPr", namespaces=NS), catalog))
        self.assertEqual(paragraph.xpath("w:pPr/w:outlineLvl/@w:val", namespaces=NS), ["9"])
        self.assertEqual(stats.heading_numbers_applied, 0)

    def test_outline_demotion_happens_once_and_headers_keep_their_level(self) -> None:
        target = self.working_dir / "outline-scope-target.docx"
        output = self.working_dir / "outline-scope-output.docx"
        document = Document()
        body = document.add_paragraph("正文大纲二级")
        _set_direct_outline(body, 1)
        deepest = document.add_paragraph("大纲九级边界")
        _set_direct_outline(deepest, 8)
        table = document.add_table(rows=1, cols=1)
        table.style = "Table Grid"
        table_paragraph = table.cell(0, 0).paragraphs[0]
        table_paragraph.text = "表格大纲二级"
        _set_direct_outline(table_paragraph, 1)
        table_paragraph.paragraph_format.space_before = Pt(13)
        header = document.sections[0].header.paragraphs[0]
        header.text = "页眉大纲二级"
        _set_direct_outline(header, 1)
        document.save(target)

        _manifest, stats = manager.apply_style_pack(
            self.pack, target, output, demote_headings=True
        )
        entries = read_zip(output)
        root = xml(entries, "word/document.xml")
        for text in ("正文大纲二级", "表格大纲二级"):
            paragraph = paragraph_for_text(root, text)
            self.assertEqual(paragraph_style(paragraph), "Heading3")
            self.assertEqual(direct_numbering(paragraph)[1], 2)
            self.assertIsNone(paragraph.find("w:pPr/w:outlineLvl", namespaces=NS))
        table_paragraph = paragraph_for_text(root, "表格大纲二级")
        self.assertEqual(
            table_paragraph.find("w:pPr/w:spacing", namespaces=NS).get(qn("before")),
            "260",
        )
        self.assertEqual(len(root.xpath("//w:tbl", namespaces=NS)), 1)
        original_root = xml(read_zip(target), "word/document.xml")
        self.assertEqual(
            [etree.tostring(node, method="c14n", exclusive=True) for node in root.xpath("//w:tblPr", namespaces=NS)],
            [etree.tostring(node, method="c14n", exclusive=True) for node in original_root.xpath("//w:tblPr", namespaces=NS)],
        )
        deepest_paragraph = paragraph_for_text(root, "大纲九级边界")
        self.assertEqual(paragraph_style(deepest_paragraph), "Heading9")
        self.assertEqual(direct_numbering(deepest_paragraph)[1], 8)
        header_name = next(name for name in entries if name.startswith("word/header") and name.endswith(".xml"))
        header_paragraph = paragraph_for_text(xml(entries, header_name), "页眉大纲二级")
        self.assertEqual(paragraph_style(header_paragraph), "Heading2")
        self.assertEqual(direct_numbering(header_paragraph)[1], 1)
        self.assertEqual(stats.heading_levels_demoted, 2)
        self.assertEqual(stats.heading_level9_unchanged, 1)

    def test_direct_outline_uses_custom_source_heading_authority(self) -> None:
        target = self.working_dir / "custom-authority-outline-target.docx"
        output = self.working_dir / "custom-authority-outline-output.docx"
        pack = self.working_dir / "custom-outline.wfstyle"
        document = Document()
        _set_direct_outline(document.add_paragraph("自定义模板大纲标题"), 0)
        document.save(target)
        manager.create_style_pack(TEST_DIR / "fixtures" / "source-custom-outline-numbered.docx", pack)

        source = TEST_DIR / "fixtures" / "source-custom-outline-numbered.docx"
        for method in ("pack", "direct"):
            with self.subTest(method=method):
                output = self.working_dir / ("custom-outline-%s.docx" % method)
                if method == "pack":
                    manager.apply_style_pack(pack, target, output)
                else:
                    core.transfer(source, target, output)
                paragraph = paragraph_for_text(
                    xml(read_zip(output), "word/document.xml"),
                    "自定义模板大纲标题",
                )
                self.assertEqual(paragraph_style(paragraph), "CustomOutlineOne")
                self.assertEqual(direct_numbering(paragraph)[1], 0)

    def test_direct_transfer_outline_matrix_demotion_and_story_scope(self) -> None:
        target = self.working_dir / "direct-outline-matrix.docx"
        document = Document()
        expected = []
        for level in range(9):
            for explicit_normal in (False, True):
                text = "直接迁移大纲%d-%s" % (level, explicit_normal)
                paragraph = document.add_paragraph(text)
                if explicit_normal:
                    _set_explicit_normal_style(paragraph)
                _set_direct_outline(paragraph, level)
                expected.append((text, level))
        invalid_styles = []
        for value in (-1, 10, "invalid", ""):
            body = document.add_paragraph("非法正文大纲-%s" % value)
            _set_direct_outline(body, value)
            inherited = document.add_paragraph(
                "非法大纲继承标题-%s" % value, style="Heading 1"
            )
            _set_direct_outline(inherited, value)
            invalid_styles.append(body.text)
            expected.append((inherited.text, 0))
        for style_name in (None, "Normal", "Heading 1"):
            body = document.add_paragraph("直接正文9-%s" % style_name, style=style_name)
            if style_name == "Normal":
                _set_explicit_normal_style(body)
            _set_direct_outline(body, 9)
            invalid_styles.append(body.text)
        table = document.add_table(rows=1, cols=1)
        table.style = "Table Grid"
        cell = table.cell(0, 0).paragraphs[0]
        cell.text = "直接表格大纲二级"
        _set_explicit_normal_style(cell)
        _set_direct_outline(cell, 1)
        cell.paragraph_format.space_before = Pt(13)
        expected.append((cell.text, 1))
        header = document.sections[0].header.paragraphs[0]
        header.text = "直接页眉大纲二级"
        _set_explicit_normal_style(header)
        _set_direct_outline(header, 1)
        document.save(target)
        original_entries = read_zip(target)
        original_root = xml(original_entries, "word/document.xml")
        original_text = original_root.xpath("//w:t/text()", namespaces=NS)
        original_table_properties = [
            etree.tostring(node, method="c14n", exclusive=True)
            for node in original_root.xpath("//w:tblPr", namespaces=NS)
        ]

        for demote in (False, True):
            with self.subTest(demote=demote):
                output = self.working_dir / ("direct-outline-matrix-%s.docx" % demote)
                stats = core.transfer(self.source, target, output, demote_headings=demote)
                entries = read_zip(output)
                root = xml(entries, "word/document.xml")
                catalog = core.build_style_catalog(entries["word/styles.xml"], entries["word/document.xml"])
                self.assertEqual(root.xpath("//w:t/text()", namespaces=NS), original_text)
                for text, level in expected:
                    paragraph = paragraph_for_text(root, text)
                    mapped_level = min(8, level + int(demote))
                    self.assertEqual(paragraph_style(paragraph), "Heading%d" % (mapped_level + 1))
                    self.assertEqual(direct_numbering(paragraph)[1], mapped_level)
                    self.assertIsNone(paragraph.find("w:pPr/w:outlineLvl", namespaces=NS))
                for text in invalid_styles:
                    paragraph = paragraph_for_text(root, text)
                    style = paragraph.find("w:pPr/w:pStyle", namespaces=NS)
                    style_id = str(style.get(qn("val"))) if style is not None else catalog.fallback("paragraph")
                    self.assertIsNone(catalog.resolved_outline.get(style_id))
                    self.assertIsNone(paragraph.find("w:pPr/w:numPr", namespaces=NS))
                header_name = next(name for name in entries if name.startswith("word/header") and name.endswith(".xml"))
                header_paragraph = paragraph_for_text(xml(entries, header_name), header.text)
                self.assertEqual(paragraph_style(header_paragraph), "Heading2")
                self.assertEqual(direct_numbering(header_paragraph)[1], 1)
                self.assertEqual(
                    [etree.tostring(node, method="c14n", exclusive=True) for node in root.xpath("//w:tblPr", namespaces=NS)],
                    original_table_properties,
                )
                self.assertEqual(
                    paragraph_for_text(root, cell.text).find("w:pPr/w:spacing", namespaces=NS).get(qn("before")),
                    "260",
                )
                self.assertEqual(stats.heading_levels_demoted, 21 if demote else 0)
                self.assertEqual(stats.heading_level9_unchanged, 2 if demote else 0)

    def test_direct_outline_demotion_completes_missing_source_level(self) -> None:
        source = self.working_dir / "outline-three-level-source.docx"
        pack = self.working_dir / "outline-three-level.wfstyle"
        target = self.working_dir / "outline-five-target.docx"
        output = self.working_dir / "outline-five-demoted.docx"
        _make_three_level_numbered_source(source)
        manager.create_style_pack(source, pack)
        document = Document()
        _set_direct_outline(document.add_paragraph("大纲五级需补全六级"), 4)
        document.save(target)

        _manifest, stats = manager.apply_style_pack(pack, target, output, demote_headings=True)
        entries = read_zip(output)
        paragraph = paragraph_for_text(xml(entries, "word/document.xml"), "大纲五级需补全六级")
        self.assertEqual(paragraph_style(paragraph), "Heading6")
        self.assertEqual(direct_numbering(paragraph)[1], 5)
        self.assertEqual(stats.heading_levels_demoted, 1)

        direct_output = self.working_dir / "outline-five-direct-demoted.docx"
        direct_stats = core.transfer(source, target, direct_output, demote_headings=True)
        direct_paragraph = paragraph_for_text(
            xml(read_zip(direct_output), "word/document.xml"), "大纲五级需补全六级"
        )
        self.assertEqual(paragraph_style(direct_paragraph), "Heading6")
        self.assertEqual(direct_numbering(direct_paragraph)[1], 5)
        self.assertEqual(direct_stats.heading_levels_demoted, 1)


if __name__ == "__main__":
    unittest.main(verbosity=2)
