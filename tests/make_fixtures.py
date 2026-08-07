#!/usr/bin/env python3
from pathlib import Path
import zipfile

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor
from lxml import etree
from PIL import Image, ImageDraw


ROOT = Path(__file__).resolve().parent
FIXTURES = ROOT / "fixtures"
FIXTURES.mkdir(parents=True, exist_ok=True)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_table_borders(table, color):
    tbl_pr = table._tbl.tblPr
    existing = tbl_pr.find(qn("w:tblBorders"))
    if existing is not None:
        tbl_pr.remove(existing)
    borders = OxmlElement("w:tblBorders")
    for edge_name in ("top", "left", "bottom", "right", "insideH", "insideV"):
        edge = OxmlElement("w:%s" % edge_name)
        edge.set(qn("w:val"), "single")
        edge.set(qn("w:sz"), "12")
        edge.set(qn("w:space"), "0")
        edge.set(qn("w:color"), color)
        borders.append(edge)
    tbl_pr.append(borders)


def set_table_cell_margins(table, value):
    tbl_pr = table._tbl.tblPr
    existing = tbl_pr.find(qn("w:tblCellMar"))
    if existing is not None:
        tbl_pr.remove(existing)
    margins = OxmlElement("w:tblCellMar")
    for edge_name in ("top", "left", "bottom", "right"):
        edge = OxmlElement("w:%s" % edge_name)
        edge.set(qn("w:w"), str(value))
        edge.set(qn("w:type"), "dxa")
        margins.append(edge)
    tbl_pr.append(margins)


def set_cell_margins(cell, value):
    tc_pr = cell._tc.get_or_add_tcPr()
    existing = tc_pr.find(qn("w:tcMar"))
    if existing is not None:
        tc_pr.remove(existing)
    margins = OxmlElement("w:tcMar")
    for edge_name in ("top", "left", "bottom", "right"):
        edge = OxmlElement("w:%s" % edge_name)
        edge.set(qn("w:w"), str(value))
        edge.set(qn("w:type"), "dxa")
        margins.append(edge)
    tc_pr.append(margins)


def set_row_height(row, value):
    tr_pr = row._tr.get_or_add_trPr()
    height = tr_pr.find(qn("w:trHeight"))
    if height is None:
        height = OxmlElement("w:trHeight")
        tr_pr.append(height)
    height.set(qn("w:val"), str(value))
    height.set(qn("w:hRule"), "atLeast")


def set_default_table_style(doc, style_id):
    settings = doc.settings._element
    current = settings.find(qn("w:defaultTableStyle"))
    if current is None:
        current = OxmlElement("w:defaultTableStyle")
        default_tab_stop = settings.find(qn("w:defaultTabStop"))
        insertion = (
            settings.index(default_tab_stop)
            if default_tab_stop is not None
            else len(settings)
        )
        settings.insert(insertion, current)
    current.set(qn("w:val"), style_id)


def append_styles_extension(doc):
    styles = doc.styles._element
    for child in list(styles):
        if child.tag == qn("w:extLst"):
            styles.remove(child)
    extension_list = OxmlElement("w:extLst")
    extension = OxmlElement("w:ext")
    extension.set("uri", "{1B4A2E4E-7E4C-4D6C-8E0D-6E662AA10F35}")
    extension_list.append(extension)
    styles.append(extension_list)


def set_paragraph_border(paragraph, color):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "18")
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)
    p_pr.append(p_bdr)


def set_style_font(style, latin_name, east_asian_name):
    style.font.name = latin_name
    r_pr = style._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, r_fonts)
    r_fonts.set(qn("w:ascii"), latin_name)
    r_fonts.set(qn("w:hAnsi"), latin_name)
    r_fonts.set(qn("w:eastAsia"), east_asian_name)


def set_style_outline_level(style, level):
    p_pr = style._element.get_or_add_pPr()
    existing = p_pr.find(qn("w:outlineLvl"))
    if existing is not None:
        p_pr.remove(existing)
    outline = OxmlElement("w:outlineLvl")
    outline.set(qn("w:val"), str(level))
    p_pr.append(outline)


def _next_numbering_id(root, attribute_name):
    values = []
    for node in root:
        value = node.get(qn(attribute_name))
        if value is None:
            continue
        try:
            values.append(int(value))
        except ValueError:
            continue
    return max(values, default=-1) + 1


def _value_child(tag, value):
    child = OxmlElement(tag)
    child.set(qn("w:val"), str(value))
    return child


def set_direct_paragraph_numbering(paragraph, num_id, level=None):
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = p_pr.get_or_add_numPr()
    if level is None:
        if num_pr.ilvl is not None:
            num_pr._remove_ilvl()
    else:
        num_pr.get_or_add_ilvl().val = level
    num_pr.get_or_add_numId().val = num_id


def set_style_numbering(style, num_id, level):
    num_pr = style._element.get_or_add_pPr().get_or_add_numPr()
    num_pr.get_or_add_ilvl().val = level
    num_pr.get_or_add_numId().val = num_id


def add_heading_multilevel_numbering(
    doc, attach_to_styles=True, link_level_styles=True
):
    """Attach one real 3-level Word numbering definition to Heading 1/2/3."""
    numbering = doc.part.numbering_part.element
    abstract_id = _next_numbering_id(numbering, "w:abstractNumId")
    num_id = _next_numbering_id(numbering, "w:numId")

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    abstract.append(_value_child("w:nsid", "A11CE123"))
    abstract.append(_value_child("w:multiLevelType", "multilevel"))
    abstract.append(_value_child("w:tmpl", "C0DE1234"))

    level_specs = (
        (0, "Heading1", "%1", 432),
        (1, "Heading2", "%1.%2", 576),
        (2, "Heading3", "%1.%2.%3", 720),
    )
    for level, style_id, level_text, indent in level_specs:
        lvl = OxmlElement("w:lvl")
        lvl.set(qn("w:ilvl"), str(level))
        lvl.append(_value_child("w:start", 1))
        lvl.append(_value_child("w:numFmt", "decimal"))
        if link_level_styles:
            lvl.append(_value_child("w:pStyle", style_id))
        lvl.append(_value_child("w:suff", "tab"))
        lvl.append(_value_child("w:lvlText", level_text))
        lvl.append(_value_child("w:lvlJc", "left"))

        p_pr = OxmlElement("w:pPr")
        tabs = OxmlElement("w:tabs")
        tab = OxmlElement("w:tab")
        tab.set(qn("w:val"), "num")
        tab.set(qn("w:pos"), str(indent))
        tabs.append(tab)
        p_pr.append(tabs)
        indentation = OxmlElement("w:ind")
        indentation.set(qn("w:left"), str(indent))
        indentation.set(qn("w:hanging"), str(indent))
        p_pr.append(indentation)
        lvl.append(p_pr)
        abstract.append(lvl)

    # CT_Numbering requires abstractNum elements before concrete num elements.
    insertion = len(numbering)
    for index, child in enumerate(numbering):
        if child.tag == qn("w:num"):
            insertion = index
            break
    numbering.insert(insertion, abstract)

    concrete = OxmlElement("w:num")
    concrete.set(qn("w:numId"), str(num_id))
    concrete.append(_value_child("w:abstractNumId", abstract_id))
    numbering.append(concrete)

    if attach_to_styles:
        for level, style_name in enumerate(
            ("Heading 1", "Heading 2", "Heading 3")
        ):
            set_style_numbering(doc.styles[style_name], num_id, level)
    return num_id


def add_single_level_numbering(doc, style, style_id, level_text="%1"):
    """Create a concrete one-level definition and attach it to one style."""
    numbering = doc.part.numbering_part.element
    abstract_id = _next_numbering_id(numbering, "w:abstractNumId")
    num_id = _next_numbering_id(numbering, "w:numId")

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    abstract.append(_value_child("w:nsid", "A11CE456"))
    abstract.append(_value_child("w:multiLevelType", "singleLevel"))
    abstract.append(_value_child("w:tmpl", "C0DE5678"))
    lvl = OxmlElement("w:lvl")
    lvl.set(qn("w:ilvl"), "0")
    lvl.append(_value_child("w:start", 1))
    lvl.append(_value_child("w:numFmt", "decimal"))
    lvl.append(_value_child("w:pStyle", style_id))
    lvl.append(_value_child("w:suff", "tab"))
    lvl.append(_value_child("w:lvlText", level_text))
    lvl.append(_value_child("w:lvlJc", "left"))
    p_pr = OxmlElement("w:pPr")
    indentation = OxmlElement("w:ind")
    indentation.set(qn("w:left"), "432")
    indentation.set(qn("w:hanging"), "432")
    p_pr.append(indentation)
    lvl.append(p_pr)
    abstract.append(lvl)

    insertion = len(numbering)
    for index, child in enumerate(numbering):
        if child.tag == qn("w:num"):
            insertion = index
            break
    numbering.insert(insertion, abstract)

    concrete = OxmlElement("w:num")
    concrete.set(qn("w:numId"), str(num_id))
    concrete.append(_value_child("w:abstractNumId", abstract_id))
    numbering.append(concrete)
    set_style_numbering(style, num_id, 0)
    return num_id


def add_numbering_style_proxy(doc, actual_num_id, target_style):
    """Create a numStyleLink proxy and point target_style at its numId."""
    proxy_style = doc.styles.add_style(
        "Heading Number Proxy", WD_STYLE_TYPE.LIST
    )
    assert proxy_style.style_id == "HeadingNumberProxy"
    set_style_numbering(proxy_style, actual_num_id, 0)

    numbering = doc.part.numbering_part.element
    abstract_id = _next_numbering_id(numbering, "w:abstractNumId")
    proxy_num_id = _next_numbering_id(numbering, "w:numId")
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    abstract.append(_value_child("w:nsid", "A11CE789"))
    abstract.append(_value_child("w:multiLevelType", "singleLevel"))
    abstract.append(_value_child("w:numStyleLink", proxy_style.style_id))
    insertion = len(numbering)
    for index, child in enumerate(numbering):
        if child.tag == qn("w:num"):
            insertion = index
            break
    numbering.insert(insertion, abstract)

    concrete = OxmlElement("w:num")
    concrete.set(qn("w:numId"), str(proxy_num_id))
    concrete.append(_value_child("w:abstractNumId", abstract_id))
    numbering.append(concrete)
    set_style_numbering(target_style, proxy_num_id, 0)
    return proxy_num_id, proxy_style.style_id


def _set_effects_style_numbering(styles_root, style_id, num_id, level):
    style = styles_root.find(
        "w:style[@w:styleId='%s']" % style_id,
        namespaces={"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"},
    )
    if style is None:
        raise AssertionError("missing stylesWithEffects style: %s" % style_id)
    p_pr = style.find(qn("w:pPr"))
    if p_pr is None:
        p_pr = etree.Element(qn("w:pPr"))
        style.append(p_pr)
    for existing in p_pr.findall(qn("w:numPr")):
        p_pr.remove(existing)

    num_pr = etree.Element(qn("w:numPr"))
    ilvl = etree.SubElement(num_pr, qn("w:ilvl"))
    ilvl.set(qn("w:val"), str(level))
    concrete = etree.SubElement(num_pr, qn("w:numId"))
    concrete.set(qn("w:val"), str(num_id))

    # Keep the pPr child order schema-valid: numPr precedes spacing/ind/outlineLvl.
    insert_before = {
        "suppressLineNumbers",
        "pBdr",
        "shd",
        "tabs",
        "spacing",
        "ind",
        "contextualSpacing",
        "mirrorIndents",
        "suppressOverlap",
        "jc",
        "textDirection",
        "textAlignment",
        "textboxTightWrap",
        "outlineLvl",
        "divId",
        "cnfStyle",
    }
    insertion = len(p_pr)
    for index, child in enumerate(p_pr):
        if etree.QName(child).localname in insert_before:
            insertion = index
            break
    p_pr.insert(insertion, num_pr)


def mirror_numbering_to_styles_with_effects(path, assignments):
    """Keep Word's modern stylesWithEffects part aligned with styles.xml."""
    with zipfile.ZipFile(path, "r") as archive:
        infos = archive.infolist()
        entries = {info.filename: archive.read(info.filename) for info in infos}
    effects_name = "word/stylesWithEffects.xml"
    if effects_name not in entries:
        return
    root = etree.fromstring(entries[effects_name])
    for style_id, num_id, level in assignments:
        _set_effects_style_numbering(root, style_id, num_id, level)
    entries[effects_name] = etree.tostring(
        root, xml_declaration=True, encoding="UTF-8", standalone=True
    )

    temporary = path.with_suffix(path.suffix + ".tmp")
    with zipfile.ZipFile(temporary, "w") as archive:
        for info in infos:
            archive.writestr(info, entries[info.filename])
    temporary.replace(path)


def mirror_heading_numbering_to_styles_with_effects(path, num_id):
    mirror_numbering_to_styles_with_effects(
        path,
        tuple(
            (style_id, num_id, level)
            for level, style_id in enumerate(
                ("Heading1", "Heading2", "Heading3")
            )
        ),
    )


def add_hyperlink(paragraph, text, url):
    part = paragraph.part
    rel_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    r_style = OxmlElement("w:rStyle")
    r_style.set(qn("w:val"), "Hyperlink")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "FF00FF")
    bold = OxmlElement("w:b")
    r_pr.extend([r_style, color, bold])
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.extend([r_pr, text_node])
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def make_source(path: Path, include_table: bool = True):
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.4)
    section.right_margin = Cm(2.4)

    normal = doc.styles["Normal"]
    set_style_font(normal, "Arial", "Hiragino Sans GB")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor(0x22, 0x22, 0x22)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.25

    heading_specs = {
        "Heading 1": (20, RGBColor(0x16, 0x5D, 0x52), 14, 6),
        "Heading 2": (15, RGBColor(0x2F, 0x55, 0x97), 11, 4),
        "Heading 3": (12, RGBColor(0x70, 0x3A, 0x13), 8, 3),
    }
    for name, (size, color, before, after) in heading_specs.items():
        style = doc.styles[name]
        set_style_font(style, "Arial", "Hiragino Sans GB")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)

    callout = doc.styles.add_style("Source Callout", WD_STYLE_TYPE.PARAGRAPH)
    callout.base_style = normal
    set_style_font(callout, "Arial", "Hiragino Sans GB")
    callout.font.size = Pt(10)
    callout.font.italic = True
    callout.font.color.rgb = RGBColor(0x16, 0x5D, 0x52)
    callout.paragraph_format.left_indent = Cm(0.8)
    callout.paragraph_format.right_indent = Cm(0.8)
    callout.paragraph_format.space_before = Pt(6)
    callout.paragraph_format.space_after = Pt(6)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), "E8F3F1")
    callout._element.get_or_add_pPr().append(shd)

    doc.add_heading("格式源示例", level=1)
    doc.add_paragraph("这段内容只用于定义源文件的正文外观。")
    doc.add_heading("二级标题示例", level=2)
    doc.add_heading("三级标题示例", level=3)
    doc.add_paragraph("源文件提示框样式", style="Source Callout")

    if include_table:
        table = doc.add_table(rows=2, cols=2)
        table.style = "Light Shading Accent 1"
        table.cell(0, 0).text = "源表头 A"
        table.cell(0, 1).text = "源表头 B"
        table.cell(1, 0).text = "源值 1"
        table.cell(1, 1).text = "源值 2"

    header = section.header
    header.paragraphs[0].text = "源文件页眉格式示例"
    footer = section.footer
    footer.paragraphs[0].text = "源文件页脚格式示例"
    append_styles_extension(doc)
    doc.save(path)


def make_target(path: Path):
    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Cm(29.7)
    section.page_height = Cm(21)
    section.top_margin = Cm(0.8)
    section.bottom_margin = Cm(0.8)
    section.left_margin = Cm(1.0)
    section.right_margin = Cm(1.0)

    normal = doc.styles["Normal"]
    normal.font.name = "Courier New"
    normal.font.size = Pt(15)
    normal.font.color.rgb = RGBColor(0x99, 0x00, 0x00)

    for name, size in (("Heading 1", 27), ("Heading 2", 23), ("Heading 3", 19)):
        style = doc.styles[name]
        style.font.name = "Courier New"
        style.font.size = Pt(size)
        style.font.bold = False
        style.font.color.rgb = RGBColor(0xCC, 0x00, 0x00)

    target_only = doc.styles.add_style("Target Only", WD_STYLE_TYPE.PARAGRAPH)
    target_only.base_style = normal
    target_only.font.name = "Times New Roman"
    target_only.font.size = Pt(18)
    target_only.font.bold = True
    target_only.font.color.rgb = RGBColor(0xFF, 0x00, 0xFF)

    target_table = doc.styles.add_style(
        "Target Custom Table", WD_STYLE_TYPE.TABLE
    )
    target_table.base_style = doc.styles["Table Grid"]
    target_table.font.name = "Courier New"
    target_table.font.size = Pt(13)
    target_table.font.bold = True
    target_table.font.color.rgb = RGBColor(0x00, 0x55, 0x88)
    assert target_table.style_id == "TargetCustomTable"
    set_default_table_style(doc, target_table.style_id)

    p = doc.add_paragraph(style="Heading 1")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.left_indent = Cm(2)
    p.paragraph_format.space_before = Pt(30)
    set_paragraph_border(p, "FF0000")
    run = p.add_run("目标文档主标题")
    run.font.name = "Comic Sans MS"
    run.font.size = Pt(32)
    run.font.bold = True
    run.font.italic = True
    run.font.color.rgb = RGBColor(0xFF, 0x00, 0x00)

    p = doc.add_paragraph(style="Heading 2")
    p.add_run("目标文档二级标题").font.color.rgb = RGBColor(0xFF, 0x00, 0xFF)
    p = doc.add_paragraph(style="Heading 3")
    p.add_run("目标文档三级标题").font.size = Pt(25)

    p = doc.add_paragraph(style="Target Only")
    p.add_run("目标独有样式段落，输出后应映射为源文件正文。")

    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Cm(2)
    p.paragraph_format.line_spacing = 2
    run = p.add_run("正文含有手工粗体、斜体、字号和颜色。")
    run.bold = True
    run.italic = True
    run.font.size = Pt(20)
    run.font.color.rgb = RGBColor(0x00, 0x00, 0xFF)

    list_p = doc.add_paragraph("目标编号列表项目", style="List Number")
    list_p.runs[0].font.color.rgb = RGBColor(0xFF, 0x00, 0x00)

    link_p = doc.add_paragraph("外部链接：")
    add_hyperlink(link_p, "OpenAI", "https://openai.com")

    table = doc.add_table(rows=3, cols=2)
    table.style = target_table
    set_table_borders(table, "0088CC")
    set_table_cell_margins(table, 180)
    values = (("目标表头 A", "目标表头 B"), ("目标值 1", "目标值 2"), ("目标值 3", "目标值 4"))
    for row_index, row_values in enumerate(values):
        for col_index, value in enumerate(row_values):
            cell = table.cell(row_index, col_index)
            cell.text = value
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.BOTTOM
            set_cell_shading(cell, "FFFF00" if row_index == 0 else "FFCCCC")
            cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(0xFF, 0x00, 0x00)
    set_row_height(table.rows[0], 620)
    set_cell_margins(table.cell(2, 0), 240)
    table.cell(1, 0).merge(table.cell(1, 1))

    image_path = FIXTURES / "target-image.png"
    image = Image.new("RGB", (360, 110), "#e8f3f1")
    draw = ImageDraw.Draw(image)
    draw.rectangle((4, 4, 355, 105), outline="#165d52", width=4)
    draw.text((22, 40), "TARGET IMAGE CONTENT", fill="#165d52")
    image.save(image_path)
    doc.add_picture(str(image_path), width=Inches(3.2))

    header = section.header
    header_p = header.paragraphs[0]
    header_p.text = "必须保留的目标页眉内容"
    header_p.runs[0].font.size = Pt(22)
    header_p.runs[0].font.color.rgb = RGBColor(0xFF, 0x00, 0x00)
    footer = section.footer
    footer_p = footer.paragraphs[0]
    footer_p.text = "必须保留的目标页脚内容"
    footer_p.runs[0].bold = True
    footer_p.runs[0].font.color.rgb = RGBColor(0xFF, 0x00, 0xFF)

    doc.save(path)


def format_heading_source_styles(doc):
    normal = doc.styles["Normal"]
    set_style_font(normal, "Arial", "Hiragino Sans GB")
    normal.font.size = Pt(11)

    heading_specs = (
        ("Heading 1", 20, RGBColor(0x16, 0x5D, 0x52)),
        ("Heading 2", 15, RGBColor(0x2F, 0x55, 0x97)),
        ("Heading 3", 12, RGBColor(0x70, 0x3A, 0x13)),
    )
    for name, size, color in heading_specs:
        style = doc.styles[name]
        set_style_font(style, "Arial", "Hiragino Sans GB")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color


def make_three_level_heading_logic_source(path: Path, numbered: bool):
    """Create a source whose used H1-H3 form a clear design progression.

    Heading 4/5 remain defined in the package but are deliberately corrupted
    and unused.  A correct style-learning implementation must extrapolate from
    the three actually used heading levels instead of copying those dormant
    definitions.
    """
    doc = Document()
    normal = doc.styles["Normal"]
    set_style_font(normal, "Arial", "PingFang SC")
    normal.font.size = Pt(10)
    normal.font.color.rgb = RGBColor(0x22, 0x22, 0x22)
    normal.paragraph_format.space_after = Pt(5)

    # These values intentionally follow constant deltas so inference is
    # deterministic: size -3pt, before -4pt, after -2pt per level.
    heading_specs = (
        ("Heading 1", 21, 18, 8),
        ("Heading 2", 18, 14, 6),
        ("Heading 3", 15, 10, 4),
    )
    for level, (name, size, before, after) in enumerate(heading_specs):
        style = doc.styles[name]
        set_style_font(style, "Arial", "PingFang SC")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.italic = False
        style.font.color.rgb = RGBColor(0x24, 0x4A, 0x73)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        set_style_outline_level(style, level)

    bad_specs = (
        ("Heading 4", 31, 50, 30, "Comic Sans MS", "SimSun", 7, "FF00FF"),
        ("Heading 5", 29, 45, 25, "Papyrus", "KaiTi", 8, "00FF00"),
    )
    for name, size, before, after, latin, east_asian, outline, color in bad_specs:
        style = doc.styles[name]
        set_style_font(style, latin, east_asian)
        style.font.size = Pt(size)
        style.font.bold = False
        style.font.italic = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        set_style_outline_level(style, outline)

    num_id = add_heading_multilevel_numbering(doc) if numbered else None
    doc.add_heading("格式源一级标题", level=1)
    doc.add_heading("格式源二级标题", level=2)
    doc.add_heading("格式源三级标题", level=3)
    doc.add_paragraph("格式源只实际使用前三层标题。")
    doc.save(path)
    if num_id is not None:
        mirror_heading_numbering_to_styles_with_effects(path, num_id)


def make_rich_numbering_format_source(path: Path):
    """Create a 3-level source whose number labels have explicit typography.

    The three source levels deliberately use different fonts, sizes, colors,
    and emphasis.  Their paragraph geometry uses a stable 360-twip step so
    Heading 4/5 extension has a deterministic expected result.  ``suff`` and
    ``lvlJc`` are also non-default to catch lossy number-format migration.
    """
    make_three_level_heading_logic_source(path, numbered=True)

    with zipfile.ZipFile(path, "r") as archive:
        infos = archive.infolist()
        entries = {info.filename: archive.read(info.filename) for info in infos}

    styles = etree.fromstring(entries["word/styles.xml"])
    heading_one = styles.find(
        "w:style[@w:styleId='Heading1']",
        namespaces={
            "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
        },
    )
    if heading_one is None:
        raise AssertionError("missing Heading1 in rich numbering fixture")
    num_id_node = heading_one.find("w:pPr/w:numPr/w:numId", namespaces=heading_one.nsmap)
    if num_id_node is None:
        raise AssertionError("Heading1 is not numbered in rich numbering fixture")
    num_id = num_id_node.get(qn("w:val"))

    numbering = etree.fromstring(entries["word/numbering.xml"])
    concrete = numbering.find(
        "w:num[@w:numId='%s']" % num_id,
        namespaces={
            "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
        },
    )
    if concrete is None:
        raise AssertionError("missing concrete numbering definition")
    abstract_ref = concrete.find(qn("w:abstractNumId"))
    abstract_id = abstract_ref.get(qn("w:val"))
    abstract = numbering.find(
        "w:abstractNum[@w:abstractNumId='%s']" % abstract_id,
        namespaces={
            "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
        },
    )
    if abstract is None:
        raise AssertionError("missing abstract numbering definition")

    label_specs = (
        # latin, east Asia, color, half-points, emphasis
        ("Aptos Display", "STSong", "8A1538", 34, "b"),
        ("Courier New", "SimHei", "1F4E78", 30, "i"),
        # Intentionally conspicuous: blindly cloning this onto H4/H5 is wrong.
        ("Georgia", "KaiTi", "C00000", 56, "smallCaps"),
    )
    geometry = (
        (600, 300),
        (960, 300),
        (1320, 300),
    )
    for level, ((latin, east_asia, color, size, emphasis), (position, hanging)) in enumerate(
        zip(label_specs, geometry)
    ):
        lvl = abstract.find(
            "w:lvl[@w:ilvl='%d']" % level,
            namespaces={
                "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
            },
        )
        if lvl is None:
            raise AssertionError("missing source numbering level %d" % level)
        lvl.find(qn("w:suff")).set(qn("w:val"), "space")
        lvl.find(qn("w:lvlJc")).set(qn("w:val"), "right")

        tab = lvl.find("w:pPr/w:tabs/w:tab", namespaces=lvl.nsmap)
        indentation = lvl.find("w:pPr/w:ind", namespaces=lvl.nsmap)
        if tab is None or indentation is None:
            raise AssertionError("missing numbering geometry for level %d" % level)
        tab.set(qn("w:val"), "num")
        tab.set(qn("w:pos"), str(position))
        indentation.set(qn("w:left"), str(position))
        indentation.set(qn("w:hanging"), str(hanging))

        old_r_pr = lvl.find(qn("w:rPr"))
        if old_r_pr is not None:
            lvl.remove(old_r_pr)
        r_pr = OxmlElement("w:rPr")
        r_fonts = OxmlElement("w:rFonts")
        r_fonts.set(qn("w:ascii"), latin)
        r_fonts.set(qn("w:hAnsi"), latin)
        r_fonts.set(qn("w:eastAsia"), east_asia)
        r_fonts.set(qn("w:cs"), latin)
        r_pr.append(r_fonts)
        r_pr.append(_value_child("w:color", color))
        r_pr.append(_value_child("w:sz", size))
        r_pr.append(_value_child("w:szCs", size))
        r_pr.append(OxmlElement("w:%s" % emphasis))
        lvl.append(r_pr)

    entries["word/numbering.xml"] = etree.tostring(
        numbering, xml_declaration=True, encoding="UTF-8", standalone=True
    )
    temporary = path.with_suffix(path.suffix + ".tmp")
    with zipfile.ZipFile(temporary, "w") as archive:
        for info in infos:
            archive.writestr(info, entries[info.filename])
    temporary.replace(path)


def make_multi_numid_heading_source(path: Path):
    """Create compatible H1/H2/H3 rules backed by three different numIds.

    Each concrete numbering definition points at a separate abstractNum.  The
    authoritative level in each abstract keeps the intended formatting while
    non-authoritative levels are deliberately poisoned.  A correct importer
    therefore has to materialize H1, H2, and H3 from their own numIds before
    combining them.  Heading 1 starts at chapter 3 via startOverride.
    """
    make_rich_numbering_format_source(path)
    with zipfile.ZipFile(path, "r") as archive:
        infos = archive.infolist()
        entries = {info.filename: archive.read(info.filename) for info in infos}

    namespace = {
        "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    }
    styles = etree.fromstring(entries["word/styles.xml"])
    heading_one = styles.find(
        "w:style[@w:styleId='Heading1']", namespaces=namespace
    )
    original_num_id = heading_one.find(
        "w:pPr/w:numPr/w:numId", namespaces=namespace
    ).get(qn("w:val"))

    numbering = etree.fromstring(entries["word/numbering.xml"])
    original_num = numbering.find(
        "w:num[@w:numId='%s']" % original_num_id, namespaces=namespace
    )
    original_abstract_id = original_num.find(qn("w:abstractNumId")).get(
        qn("w:val")
    )
    original_abstract = numbering.find(
        "w:abstractNum[@w:abstractNumId='%s']" % original_abstract_id,
        namespaces=namespace,
    )

    # Clone the pristine source definition before giving each abstract a
    # single authoritative heading level.
    abstracts = [
        original_abstract,
        etree.fromstring(etree.tostring(original_abstract)),
        etree.fromstring(etree.tostring(original_abstract)),
    ]
    next_abstract_id = _next_numbering_id(numbering, "w:abstractNumId")
    abstract_ids = [
        original_abstract_id,
        str(next_abstract_id),
        str(next_abstract_id + 1),
    ]
    for clone, abstract_id in zip(abstracts[1:], abstract_ids[1:]):
        clone.set(qn("w:abstractNumId"), abstract_id)

    poison_colors = ("DEAD10", "DEAD20", "DEAD30")
    for authority_level, (abstract, poison_color) in enumerate(
        zip(abstracts, poison_colors)
    ):
        for level in abstract.findall("w:lvl", namespaces=namespace):
            level_index = int(level.get(qn("w:ilvl")))
            if level_index == authority_level:
                continue
            p_style = level.find("w:pStyle", namespaces=namespace)
            if p_style is not None:
                level.remove(p_style)
            r_pr = level.find("w:rPr", namespaces=namespace)
            if r_pr is not None:
                fonts = r_pr.find("w:rFonts", namespaces=namespace)
                if fonts is not None:
                    for attribute in ("ascii", "hAnsi", "eastAsia", "cs"):
                        fonts.set(qn("w:%s" % attribute), "PoisonLevel%d" % level_index)
                color = r_pr.find("w:color", namespaces=namespace)
                if color is not None:
                    color.set(qn("w:val"), poison_color)
                size = r_pr.find("w:sz", namespaces=namespace)
                if size is not None:
                    size.set(qn("w:val"), "88")
            indentation = level.find("w:pPr/w:ind", namespaces=namespace)
            tab = level.find("w:pPr/w:tabs/w:tab", namespaces=namespace)
            if indentation is not None:
                indentation.set(qn("w:left"), str(3900 + level_index * 100))
            if tab is not None:
                tab.set(qn("w:pos"), str(3900 + level_index * 100))

    # The H1 marker itself is semantic and its first visible chapter is 3.
    abstracts[0].find(
        "w:lvl[@w:ilvl='0']/w:lvlText", namespaces=namespace
    ).set(qn("w:val"), "第%1章")

    first_num = numbering.find("w:num", namespaces=namespace)
    insertion = numbering.index(first_num) if first_num is not None else len(numbering)
    for clone in abstracts[1:]:
        numbering.insert(insertion, clone)
        insertion += 1

    next_num_id = _next_numbering_id(numbering, "w:numId")
    num_ids = [original_num_id, str(next_num_id), str(next_num_id + 1)]
    for num_id, abstract_id in zip(num_ids[1:], abstract_ids[1:]):
        concrete = OxmlElement("w:num")
        concrete.set(qn("w:numId"), num_id)
        concrete.append(_value_child("w:abstractNumId", abstract_id))
        numbering.append(concrete)

    level_override = OxmlElement("w:lvlOverride")
    level_override.set(qn("w:ilvl"), "0")
    level_override.append(_value_child("w:startOverride", 3))
    original_num.append(level_override)

    assignments = tuple(
        (style_id, num_id, level)
        for level, (style_id, num_id) in enumerate(
            zip(("Heading1", "Heading2", "Heading3"), num_ids)
        )
    )
    for style_id, num_id, level in assignments:
        _set_effects_style_numbering(styles, style_id, num_id, level)
    entries["word/styles.xml"] = etree.tostring(
        styles, xml_declaration=True, encoding="UTF-8", standalone=True
    )

    effects_name = "word/stylesWithEffects.xml"
    if effects_name in entries:
        effects = etree.fromstring(entries[effects_name])
        for style_id, num_id, level in assignments:
            _set_effects_style_numbering(effects, style_id, num_id, level)
        entries[effects_name] = etree.tostring(
            effects, xml_declaration=True, encoding="UTF-8", standalone=True
        )
    entries["word/numbering.xml"] = etree.tostring(
        numbering, xml_declaration=True, encoding="UTF-8", standalone=True
    )

    temporary = path.with_suffix(path.suffix + ".tmp")
    with zipfile.ZipFile(temporary, "w") as archive:
        for info in infos:
            archive.writestr(info, entries[info.filename])
    temporary.replace(path)


def make_manual_heading_prefix_target(path: Path):
    """Create H1-H5 target content with typed H1-H3 number prefixes."""
    doc = Document()
    doc.add_paragraph("第三章  项目总览", style="Heading 1")
    doc.add_paragraph("3.1  建设范围", style="Heading 2")

    split = doc.add_paragraph(style="Heading 3")
    split.add_run("3.")
    split.add_run("1.")
    split.add_run("1")
    split.add_run("  技术路线")

    doc.add_paragraph("接口设计", style="Heading 4")
    doc.add_paragraph("字段校验", style="Heading 5")
    doc.add_paragraph("正文保留数字 3.1、3.1.1，以及第三章的历史说明。")
    doc.add_paragraph("2026 年第 3 季度数据也必须保持不变。")
    doc.save(path)


def make_five_level_heading_target(path: Path):
    """Create a target that uses H1-H5 and gives H4/H5 junk formatting."""
    doc = Document()

    for level in range(1, 6):
        style = doc.styles["Heading %d" % level]
        set_style_font(style, "Courier New", "STKaiti")
        style.font.size = Pt(38 - level)
        style.font.bold = False
        style.font.italic = True
        style.font.color.rgb = RGBColor(0xCC, 0x00, 0x00)
        style.paragraph_format.space_before = Pt(35 + level)
        style.paragraph_format.space_after = Pt(20 + level)

    headings = (
        (1, "项目总览"),
        (2, "建设范围"),
        (3, "实施路径"),
        (4, "接口设计"),
        (5, "字段校验"),
    )
    for level, text in headings:
        paragraph = doc.add_heading(text, level=level)
        if level < 4:
            continue
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.left_indent = Cm(2.6)
        paragraph.paragraph_format.right_indent = Cm(1.4)
        paragraph.paragraph_format.space_before = Pt(48)
        paragraph.paragraph_format.space_after = Pt(27)
        set_paragraph_border(paragraph, "FF0000")
        shading = OxmlElement("w:shd")
        shading.set(qn("w:fill"), "FFFF00")
        paragraph._p.get_or_add_pPr().append(shading)
        for run in paragraph.runs:
            run.font.name = "Comic Sans MS"
            run.font.size = Pt(36)
            run.font.bold = False
            run.font.italic = True
            run.font.color.rgb = RGBColor(0xFF, 0x00, 0x00)

    doc.add_paragraph("目标正文及所有标题文字必须原样保留。")
    doc.save(path)


def make_heading_numbering_source(path: Path, numbered: bool):
    doc = Document()
    format_heading_source_styles(doc)

    num_id = add_heading_multilevel_numbering(doc) if numbered else None
    doc.add_heading("格式源一级标题", level=1)
    doc.add_heading("格式源二级标题", level=2)
    doc.add_heading("格式源三级标题", level=3)
    doc.add_paragraph("格式源正文示例")
    doc.save(path)
    if num_id is not None:
        mirror_heading_numbering_to_styles_with_effects(path, num_id)


def make_direct_heading_numbering_source(path: Path):
    doc = Document()
    format_heading_source_styles(doc)
    # This is intentionally paragraph-only numbering: the styles have no
    # numPr and the abstract levels have no pStyle back-links.
    num_id = add_heading_multilevel_numbering(
        doc, attach_to_styles=False, link_level_styles=False
    )
    for level, text in enumerate(
        ("直接编号一级标题", "直接编号二级标题", "直接编号三级标题")
    ):
        paragraph = doc.add_heading(text, level=level + 1)
        set_direct_paragraph_numbering(paragraph, num_id, level)
    doc.add_paragraph("段落直接编号格式源正文。")
    doc.save(path)


def make_custom_outline_numbering_source(path: Path):
    doc = Document()
    format_heading_source_styles(doc)
    custom = doc.styles.add_style("Custom Outline One", WD_STYLE_TYPE.PARAGRAPH)
    custom.base_style = doc.styles["Normal"]
    set_style_font(custom, "Arial", "Hiragino Sans GB")
    custom.font.size = Pt(22)
    custom.font.bold = True
    custom.font.color.rgb = RGBColor(0x7A, 0x31, 0x7A)
    outline = OxmlElement("w:outlineLvl")
    outline.set(qn("w:val"), "0")
    custom._element.get_or_add_pPr().append(outline)
    assert custom.style_id == "CustomOutlineOne"
    num_id = add_single_level_numbering(
        doc, custom, custom.style_id, level_text="第%1章"
    )
    paragraph = doc.add_paragraph("自定义大纲一级标题", style=custom)
    assert paragraph.style.style_id == custom.style_id
    doc.add_paragraph("内置 Heading 1 存在于样式库中，但格式源正文没有使用它。")
    doc.save(path)
    # The custom style is new and therefore absent from the static effects
    # copy; styles.xml is the authoritative definition for this edge case.
    return num_id


def make_explicitly_cancelled_heading_source(path: Path):
    doc = Document()
    format_heading_source_styles(doc)
    num_id = add_single_level_numbering(
        doc, doc.styles["Heading 1"], "Heading1"
    )
    paragraph = doc.add_heading("显式取消编号的一级标题", level=1)
    # Word uses numId=0 on a paragraph to suppress numbering inherited from
    # its style.  No ilvl is needed for the cancellation marker.
    set_direct_paragraph_numbering(paragraph, 0, level=None)
    doc.add_paragraph("该标题在格式源里肉眼不显示序号。")
    doc.save(path)
    mirror_numbering_to_styles_with_effects(
        path, (("Heading1", num_id, 0),)
    )


def make_num_style_link_heading_source(path: Path):
    doc = Document()
    format_heading_source_styles(doc)
    actual_num_id = add_heading_multilevel_numbering(
        doc, attach_to_styles=False, link_level_styles=True
    )
    proxy_num_id, proxy_style_id = add_numbering_style_proxy(
        doc, actual_num_id, doc.styles["Heading 1"]
    )
    doc.add_heading("通过编号样式代理的一级标题", level=1)
    doc.add_paragraph("该格式源通过 numStyleLink 间接解析编号级别。")
    doc.save(path)
    # styles.xml contains the complete numbering-style proxy chain.  The
    # resulting target paragraphs also receive direct numPr, so the static
    # stylesWithEffects copy is not required to duplicate the custom proxy.
    return proxy_num_id, proxy_style_id


def make_plain_heading_target(path: Path):
    doc = Document()
    headings = (
        (1, "项目概述"),
        (2, "建设目标"),
        (3, "技术路线"),
        (2, "验收标准"),
        (1, "实施安排"),
        (2, "交付清单"),
        (3, "质量要求"),
    )
    for level, text in headings:
        paragraph = doc.add_heading(text, level=level)
        # The target deliberately contains no direct numPr and no typed prefix.
        assert paragraph._p.find("w:pPr/w:numPr", paragraph._p.nsmap) is None
    doc.add_paragraph("目标正文必须原样保留。")
    doc.save(path)


if __name__ == "__main__":
    make_source(FIXTURES / "source.docx")
    make_source(FIXTURES / "source-no-table.docx", include_table=False)
    make_target(FIXTURES / "target.docx")
    make_heading_numbering_source(
        FIXTURES / "source-numbered-headings.docx", numbered=True
    )
    make_heading_numbering_source(
        FIXTURES / "source-plain-headings.docx", numbered=False
    )
    make_direct_heading_numbering_source(
        FIXTURES / "source-direct-numbered-headings.docx"
    )
    make_custom_outline_numbering_source(
        FIXTURES / "source-custom-outline-numbered.docx"
    )
    make_explicitly_cancelled_heading_source(
        FIXTURES / "source-heading-numbering-cancelled.docx"
    )
    make_num_style_link_heading_source(
        FIXTURES / "source-num-style-link-heading.docx"
    )
    make_plain_heading_target(FIXTURES / "target-plain-headings.docx")
    print(FIXTURES)
